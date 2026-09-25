import sys
import docx
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

doc = docx.Document(r"d:\Code Full\Java\BTL_E_commerce\Ecommerce_Java\Docs\Policy\Bo_Luat_va_Chinh_Sach_San_Marketplace.docx")
for i, para in enumerate(doc.paragraphs):
    text = para.text
    if text.strip():
        print(f"[P{i}|{para.style.name}] {text}")
print("\n---TABLES---")
for ti, tbl in enumerate(doc.tables):
    print(f"\n=== TABLE {ti} ({len(tbl.rows)} rows) ===")
    for ri, row in enumerate(tbl.rows):
        cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
        print(f"  R{ri}: " + " || ".join(cells))