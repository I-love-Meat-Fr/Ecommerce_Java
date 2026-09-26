import docx
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

doc = docx.Document(r"d:\Code Full\Java\BTL_E_commerce\Ecommerce_Java\docs\BTL_CNJava_N05.docx")

for i, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else ""
    text = p.text
    if text.strip() or "Heading" in style:
        print(f"[{i}][{style}] {text}")

print("\n\n--- TABLES ---")
for ti, t in enumerate(doc.tables):
    print(f"\n=== Table {ti} ({len(t.rows)} rows x {len(t.columns)} cols) ===")
    for ri, row in enumerate(t.rows):
        cells = [c.text.replace("\n", " | ") for c in row.cells]
        print(f"  R{ri}: {cells}")
